import csv
import io
import math
import re
import statistics
from collections import Counter
from datetime import date
from typing import Any

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import create_engine, text

from app.core.config import get_settings

router = APIRouter(prefix="/pro", tags=["Data Quality, Comparison and Reports"])
settings = get_settings()

PERIOD_DAYS = {
    "1mo": 31,
    "3mo": 93,
    "6mo": 186,
    "1y": 366,
    "2y": 732,
    "5y": 1830,
}
SYMBOL_PATTERN = re.compile(r"^[A-Za-z0-9.^=_-]{1,30}$")
PRICE_FIELDS = ("open", "high", "low", "close", "volume")


def normalize_symbol(value: str) -> str:
    symbol = value.strip().upper()
    if not SYMBOL_PATTERN.fullmatch(symbol):
        raise HTTPException(status_code=400, detail=f"Invalid symbol: {value}")
    return symbol


def fetch_price_rows(symbol: str, period: str) -> list[dict[str, Any]]:
    symbol = normalize_symbol(symbol)
    if period not in PERIOD_DAYS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported period. Choose from {', '.join(PERIOD_DAYS)}.",
        )

    engine = create_engine(settings.database_url, pool_pre_ping=True)
    try:
        with engine.connect() as connection:
            rows = connection.execute(
                text("""
                    SELECT price_date, open, high, low, close, adj_close,
                           volume, source, retrieved_at
                    FROM market_prices
                    WHERE symbol = :symbol
                      AND price_date >= CURRENT_DATE - (:days * INTERVAL '1 day')
                    ORDER BY price_date ASC
                """),
                {"symbol": symbol, "days": PERIOD_DAYS[period]},
            ).mappings().all()
            return [dict(row) for row in rows]
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=503,
            detail="Could not retrieve saved prices from PostgreSQL.",
        ) from exc
    finally:
        engine.dispose()


def date_text(value: Any) -> str | None:
    if value is None:
        return None
    return value.isoformat() if hasattr(value, "isoformat") else str(value)


def calculate_metrics(rows: list[dict[str, Any]], period: str) -> dict[str, Any]:
    valid = [row for row in rows if row.get("close") is not None]
    closes = [float(row["close"]) for row in valid]
    if len(closes) < 2:
        raise ValueError("At least two valid closing prices are required.")

    returns = [
        (closes[i] / closes[i - 1]) - 1
        for i in range(1, len(closes))
        if closes[i - 1] != 0
    ]
    first, latest = closes[0], closes[-1]
    running_peak = closes[0]
    max_drawdown = 0.0

    for close in closes:
        running_peak = max(running_peak, close)
        if running_peak:
            max_drawdown = min(max_drawdown, (close / running_peak) - 1)

    changes = [closes[i] - closes[i - 1] for i in range(1, len(closes))]
    recent = changes[-14:]
    gains = [max(value, 0.0) for value in recent]
    losses = [max(-value, 0.0) for value in recent]
    avg_gain = statistics.mean(gains) if gains else 0.0
    avg_loss = statistics.mean(losses) if losses else 0.0
    rsi = (
        100.0 if avg_loss == 0 and avg_gain > 0
        else 50.0 if avg_loss == 0
        else 100.0 - (100.0 / (1.0 + avg_gain / avg_loss))
    )

    volatility = (
        statistics.stdev(returns) * math.sqrt(252) * 100
        if len(returns) > 1 else 0.0
    )

    def average(window: int) -> float | None:
        if len(closes) < window:
            return None
        return round(statistics.mean(closes[-window:]), 4)

    return {
        "period": period,
        "observations": len(closes),
        "start_date": date_text(valid[0].get("price_date")),
        "end_date": date_text(valid[-1].get("price_date")),
        "first_close": round(first, 4),
        "latest_close": round(latest, 4),
        "period_return_pct": round((latest / first - 1) * 100, 4) if first else None,
        "daily_return_pct": round(returns[-1] * 100, 4) if returns else None,
        "sma20": average(20),
        "sma50": average(50),
        "sma200": average(200),
        "rsi14_approx": round(rsi, 2),
        "annualized_volatility_pct": round(volatility, 4),
        "period_high": round(max(closes), 4),
        "period_low": round(min(closes), 4),
        "max_drawdown_pct": round(max_drawdown * 100, 4),
    }


def quality_report(
    symbol: str,
    period: str,
    rows: list[dict[str, Any]],
) -> dict[str, Any]:
    if not rows:
        raise ValueError("No saved historical prices were found for this symbol and period.")

    missing = {
        field: sum(row.get(field) is None for row in rows)
        for field in PRICE_FIELDS
    }
    dates = [row.get("price_date") for row in rows]
    duplicate_count = len(dates) - len(set(dates))
    invalid_rows = 0
    closes: list[float] = []

    for row in rows:
        close, high, low = row.get("close"), row.get("high"), row.get("low")
        bad = False
        if close is not None and float(close) <= 0:
            bad = True
        if high is not None and float(high) < 0:
            bad = True
        if low is not None and float(low) < 0:
            bad = True
        if high is not None and low is not None and float(high) < float(low):
            bad = True
        if bad:
            invalid_rows += 1
        if close is not None:
            closes.append(float(close))

    outliers = sum(
        1 for i in range(1, len(closes))
        if closes[i - 1] != 0 and abs(closes[i] / closes[i - 1] - 1) > 0.15
    )
    latest_date = dates[-1]
    if hasattr(latest_date, "date") and not isinstance(latest_date, date):
        latest_date = latest_date.date()
    stale = (
        isinstance(latest_date, date)
        and (date.today() - latest_date).days > 7
    )

    missing_cells = sum(missing.values())
    possible_cells = max(1, len(rows) * len(PRICE_FIELDS))
    score = max(
        0.0,
        min(
            100.0,
            100.0
            - (missing_cells / possible_cells * 100)
            - (invalid_rows / len(rows) * 25)
            - (outliers / max(1, len(rows) - 1) * 15)
            - (20 if duplicate_count else 0)
            - (15 if stale else 0),
        ),
    )

    return {
        "symbol": symbol,
        "period": period,
        "observations": len(rows),
        "start_date": date_text(dates[0]),
        "end_date": date_text(dates[-1]),
        "missing_counts": missing,
        "missing_cells": missing_cells,
        "duplicate_dates": duplicate_count,
        "invalid_price_rows": invalid_rows,
        "outlier_days_over_15pct": outliers,
        "data_stale": stale,
        "quality_score": round(score, 2),
        "score_note": (
            "Heuristic score based on missing fields, invalid records, "
            "large daily moves and staleness; not a statistical guarantee."
        ),
    }


@router.get("/quality/{symbol}")
def get_quality(
    symbol: str,
    period: str = Query(default="1y"),
) -> dict[str, Any]:
    rows = fetch_price_rows(symbol, period)
    try:
        return quality_report(symbol.strip().upper(), period, rows)
    except ValueError as exc:
        raise HTTPException(status_code=404, detail=str(exc)) from exc


@router.get("/compare")
def compare_stocks(
    symbols: str = Query(description="Comma-separated ticker symbols, max 8."),
    period: str = Query(default="1y"),
) -> dict[str, Any]:
    if period not in PERIOD_DAYS:
        raise HTTPException(status_code=400, detail="Unsupported period.")

    requested = list(dict.fromkeys(
        normalize_symbol(value) for value in symbols.split(",") if value.strip()
    ))
    if not requested or len(requested) > 8:
        raise HTTPException(status_code=400, detail="Provide 1 to 8 symbols.")

    results = []
    for symbol in requested:
        rows = fetch_price_rows(symbol, period)
        try:
            metrics = calculate_metrics(rows, period)
            metrics["symbol"] = symbol
            results.append(metrics)
        except ValueError:
            results.append({
                "symbol": symbol,
                "observations": len(rows),
                "unavailable": "Insufficient saved closing prices. Load the symbol history first.",
            })

    usable = [r for r in results if r.get("period_return_pct") is not None]
    if not usable:
        raise HTTPException(
            status_code=404,
            detail="No symbols have enough saved data for this comparison. Load their history first.",
        )

    return {
        "period": period,
        "results": results,
        "equal_weight_average_return_pct": round(
            statistics.mean(r["period_return_pct"] for r in usable), 4
        ),
        "note": (
            "Equal-weight average of available historical returns, not a forecast. "
            "Currency differences and FX effects are not adjusted."
        ),
    }


@router.get("/export/{symbol}")
def export_report(
    symbol: str,
    period: str = Query(default="1y"),
    export_format: str = Query(default="xlsx", alias="format"),
):
    symbol = normalize_symbol(symbol)
    rows = fetch_price_rows(symbol, period)
    if not rows:
        raise HTTPException(status_code=404, detail="No saved prices to export.")
    if export_format not in {"csv", "xlsx"}:
        raise HTTPException(status_code=400, detail="Choose format=csv or format=xlsx.")

    quality = quality_report(symbol, period, rows)
    try:
        metrics = calculate_metrics(rows, period) if sum(
            row.get("close") is not None for row in rows
        ) >= 2 else {"note": "Not enough valid closes to calculate period analytics."}
    except ValueError:
        metrics = {"note": "Not enough valid closes to calculate period analytics."}

    safe_symbol = re.sub(r"[^A-Za-z0-9._-]", "_", symbol)

    if export_format == "csv":
        output = io.StringIO()
        fields = [
            "price_date", "open", "high", "low", "close",
            "adj_close", "volume", "source", "retrieved_at",
        ]
        writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({key: date_text(row.get(key)) if key in {"price_date", "retrieved_at"} else row.get(key) for key in fields})
        content = io.BytesIO(output.getvalue().encode("utf-8-sig"))
        filename = f"MarketMind_{safe_symbol}_{period}.csv"
        return StreamingResponse(
            content,
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise HTTPException(
            status_code=503,
            detail="Excel export requires openpyxl in the backend environment.",
        ) from exc

    workbook = Workbook()
    prices_sheet = workbook.active
    prices_sheet.title = "Price History"
    fields = [
        "price_date", "open", "high", "low", "close",
        "adj_close", "volume", "source", "retrieved_at",
    ]
    prices_sheet.append(fields)
    for row in rows:
        prices_sheet.append([
            date_text(row.get(key)) if key in {"price_date", "retrieved_at"} else row.get(key)
            for key in fields
        ])

    quality_sheet = workbook.create_sheet("Data Quality")
    quality_sheet.append(["Metric", "Value"])
    for key, value in quality.items():
        if isinstance(value, dict):
            for subkey, subvalue in value.items():
                quality_sheet.append([f"{key}.{subkey}", subvalue])
        else:
            quality_sheet.append([key, value])

    summary_sheet = workbook.create_sheet("Analytics Summary")
    summary_sheet.append(["Metric", "Value"])
    for key, value in metrics.items():
        summary_sheet.append([key, value])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183D38")
        for column in sheet.columns:
            max_width = min(
                42,
                max(12, max(len(str(cell.value or "")) for cell in column) + 2),
            )
            sheet.column_dimensions[get_column_letter(column[0].column)].width = max_width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"MarketMind_{safe_symbol}_{period}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )